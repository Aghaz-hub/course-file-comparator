import streamlit as st
import pdfplumber
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re
from datetime import datetime
from pathlib import Path
import traceback
from difflib import SequenceMatcher
from io import BytesIO

st.set_page_config(
    page_title="Course File Checker & Generator",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ====================== PATHS ======================
STANDARDS_DIR = Path("standards")
STANDARDS_DIR.mkdir(exist_ok=True)
STANDARD_PATH = STANDARDS_DIR / "dr_sachin_standard"

# ====================== REQUIRED HEADINGS ======================
REQUIRED_HEADINGS = [
    "Vision and Mission of the University",
    "Vision and Mission of the School",
    "Vision and Mission of the Department",
    "GA's, PEOs, POs and PSOs of the Program",
    "Graduate Attributes",
    "Program Educational Objectives",
    "Program Outcomes",
    "Program Specific Outcomes",
    "List of KSA Profiles of the Program",
    "Knowledge, Skill and Attitude Profiles",
    "Academic Calendar",
    "Course Design of Theory component",
    "KSA Profile Elements of the Theory component",
    "List of SDGs/Ethics/Skill Levels",
    "Course Objective and List of Course Outcomes",
    "Mapping to Skill/Employment/Entrepreneurship",
    "Relevance of the Theory component",
    "Syllabus with Text books, Reference books, Online Resources, MOOCs",
    "Lesson Plan",
    "CO-PO Mapping",
    "CO-KSA Mapping",
    "Course Assessment Plan",
    "Course Design of Lab component",
    "List of Experiments",
    "List of Miniprojects",
    "Setting CO Targets",
    "List of faculty members teaching the course",
    "Time table of concerned faculty members",
    "Lecture Notes and PowerPoint handouts",
    "Tutorials/Assignments",
    "Solution of Tutorials/Assignments",
    "Mid-Sem Question Paper",
    "End-Term Question Paper",
    "Solution of Mid-Sem and End-Term Question Papers",
    "Question Papers / Continuous Assessment",
    "Rubrics for Lab assessments",
    "Note on how the course can contribute to the ePortfolio",
    "Attendance Record",
    "Continuous Evaluation for all Assessments",
    "Assessment-wise Marks Uploaded on ERP",
    "Records of Rubrics based Evaluation",
    "Minutes of Course Coordinator Level Meetings",
    "CO attainment analysis and action taken",
    "List of slow Learners and Advanced Learners",
    "The action taken to support slow and advanced learners",
    "Impact Analysis Report and Action Taken",
    "Alternative Plan of Action",
    "Sample hardcopies of Tutorials",
    "Sample hardcopies of answer sheets",
]

# ====================== HELPERS ======================
def save_uploaded_file(uploaded_file, save_path_without_ext):
    if uploaded_file is None:
        return None
    ext = Path(uploaded_file.name).suffix.lower()
    full_path = Path(str(save_path_without_ext) + ext)
    for old_ext in [".pdf", ".docx", ".txt"]:
        old = Path(str(save_path_without_ext) + old_ext)
        if old.exists() and old != full_path:
            try:
                old.unlink()
            except Exception:
                pass
    with open(full_path, "wb") as f:
        f.write(uploaded_file.getbuffer())
    return full_path

def load_standard_file():
    for ext in [".pdf", ".docx", ".txt"]:
        path = Path(str(STANDARD_PATH) + ext)
        if path.exists():
            return path
    return None

def extract_full_text(source, is_upload=False):
    try:
        if is_upload:
            name = source.name.lower()
            file_obj = source
        else:
            name = source.name.lower()
            file_obj = source

        if name.endswith(".pdf"):
            text = ""
            with pdfplumber.open(file_obj) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text += t + "\n"
            return text
        elif name.endswith(".docx"):
            doc = Document(file_obj)
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        elif name.endswith(".txt"):
            if is_upload:
                return file_obj.read().decode("utf-8", errors="ignore")
            else:
                return Path(file_obj).read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        st.warning(f"Text extraction error: {e}")
    return ""

def normalize(text):
    text = text.lower()
    text = re.sub(r"[^\w\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text

def similarity(a, b):
    return SequenceMatcher(None, normalize(a), normalize(b)).ratio()

def find_heading_in_text(required_heading, full_text, threshold=0.65):
    lines = [l.strip() for l in full_text.split("\n") if l.strip()]
    best_score = 0.0
    best_line = ""
    best_idx = -1

    for i, line in enumerate(lines):
        if len(line) > 180:
            continue
        score = similarity(required_heading, line)
        if score > best_score:
            best_score = score
            best_line = line
            best_idx = i

    found = best_score >= threshold
    content_length = 0
    if found and best_idx >= 0:
        content_lines = []
        for j in range(best_idx + 1, min(best_idx + 12, len(lines))):
            content_lines.append(lines[j])
        content_length = sum(len(c) for c in content_lines)

    return {
        "found": found,
        "matched_line": best_line,
        "score": best_score,
        "content_length": content_length
    }

def check_required_headings(full_text, threshold=0.65, min_content=30):
    results = []
    missing_count = 0
    for req in REQUIRED_HEADINGS:
        res = find_heading_in_text(req, full_text, threshold)
        status = "Present"
        reason = ""
        if not res["found"]:
            status = "Missing"
            reason = "Heading not found"
            missing_count += 1
        elif res["content_length"] < min_content:
            status = "Blank / Incomplete"
            reason = f"Heading found but content is too short ({res['content_length']} chars)"
            missing_count += 1
        else:
            reason = "OK"
        results.append({
            "required": req,
            "status": status,
            "reason": reason,
            "matched_line": res["matched_line"],
            "score": res["score"],
            "content_length": res["content_length"]
        })
    total = len(REQUIRED_HEADINGS)
    coverage = round(((total - missing_count) / total) * 100, 1)
    return results, coverage, missing_count

def generate_report(results, coverage, missing_count, your_name, std_name, threshold):
    report = f"""# Course File Structure Checklist Report
**Generated:** {datetime.now().strftime("%Y-%m-%d %H:%M")}
**Faculty File:** {your_name}
**Standard Template:** {std_name}
**Heading Match Threshold:** {threshold}
**Coverage Score:** {coverage}%
**Missing / Blank sections:** {missing_count}

---

## Detailed Checklist

| # | Required Heading | Status | Remarks |
|---|------------------|--------|---------|
"""
    for i, r in enumerate(results, 1):
        report += f"| {i} | {r['required']} | {r['status']} | {r['reason']} |\n"

    report += "\n---\n\n## Missing / Incomplete Sections\n\n"
    miss = [r for r in results if r["status"] != "Present"]
    if not miss:
        report += "All required structural headings are present and have content.\n"
    else:
        for i, r in enumerate(miss, 1):
            report += f"{i}. **{r['required']}**\n"
            report += f"   - Status: {r['status']}\n"
            report += f"   - Reason: {r['reason']}\n"
            if r["matched_line"]:
                report += f"   - Closest match found: `{r['matched_line']}` (score {r['score']:.2f})\n"
            report += "\n"
    return report

# ====================== COURSE FILE GENERATOR ======================
def set_run_font(run, size=11, bold=False):
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    except Exception:
        pass
    run.bold = bold

def add_heading_para(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 0:
        set_run_font(run, 16, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        set_run_font(run, 14, bold=True)
    else:
        set_run_font(run, 12, bold=True)
    p.space_after = Pt(6)
    return p

def add_body(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 11)
    p.space_after = Pt(4)
    return p

def add_placeholder(doc, text="[To be filled by faculty]"):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 11)
    run.italic = True
    try:
        run.font.color.rgb = RGBColor(128, 128, 128)
    except Exception:
        pass
    return p

def create_course_file(course_code, course_name, faculty_name, program, session, syllabus_text, ltpc="3-0-2-4"):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Cover
    add_heading_para(doc, "SCHOOL OF ENGINEERING", 0)
    add_heading_para(doc, "DEPARTMENT OF COMPUTER SCIENCE & TECHNOLOGY", 0)
    doc.add_paragraph()
    add_heading_para(doc, "COURSE FILE", 0)
    doc.add_paragraph()
    add_heading_para(doc, course_name.upper(), 0)
    add_heading_para(doc, f"({course_code})", 0)
    doc.add_paragraph()

    add_body(doc, f"Program(s) and semester: {program}")
    add_body(doc, "NCrF Level of Program: Level 6")
    add_body(doc, f"Course Faculty: {faculty_name}")
    add_body(doc, f"Session: {session}")
    doc.add_paragraph()

    # Checklist
    add_heading_para(doc, "COURSE FILE CONTENTS", 1)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "S. No."
    hdr[1].text = "Description"
    hdr[2].text = "Y/N"

    checklist_items = [
        ("OUTCOME BASED CURRICULUM", True),
        ("1. Vision and Mission of the University", False),
        ("2. Vision and Mission of the School", False),
        ("3. Vision and Mission of the Department, if applicable", False),
        ("4. GA's, PEOs, POs and PSOs of the Program", False),
        ("5. List of KSA Profiles of the Program", False),
        ("6. Academic Calendar", False),
        ("OUTCOME BASED COURSE DESIGN", True),
        ("7. Course Design of Theory component", False),
        ("8. Course Design of Lab component", False),
        ("OUTCOME BASED TEACHING & LEARNING", True),
        ("9. Setting CO Targets", False),
        ("10. List of faculty members teaching the course", False),
        ("11. Time table of concerned faculty members", False),
        ("12. Lecture Notes and PowerPoint handouts", False),
        ("13. Tutorials/Assignments", False),
        ("14. Solution of Tutorials/Assignments", False),
        ("15. Mid-Sem Question Paper", False),
        ("16. End-Term Question Paper", False),
        ("17. Solution of Mid-Sem and End-Term Question Papers", False),
        ("18. Question Papers / Continuous Assessment", False),
        ("19. Rubrics for Lab assessments", False),
        ("20. Note on ePortfolio of Students", False),
        ("21. Attendance Record for theory and labs", False),
        ("OUTCOME-BASED ASSESSMENT & ATTAINMENT", True),
        ("22. Continuous Evaluation from ERP", False),
        ("23. Assessment-wise Marks Uploaded on ERP", False),
        ("24. Records of Rubrics based Evaluation", False),
        ("25. Records of Assessments based on ePortfolios", False),
        ("26. Minutes of Course Coordinator Level Meetings", False),
        ("27. CO attainment analysis and action taken", False),
        ("28. List of slow Learners and Advanced Learners", False),
        ("29. Action taken to support slow and advanced learners", False),
        ("30. Impact Analysis Report and Action Taken", False),
        ("31. Alternative Plan of Action", False),
        ("APPENDICES", True),
        ("A. Sample hardcopies of Tutorials and Assignments", False),
        ("B. Sample hardcopies of answer sheets", False),
    ]

    for item, is_header in checklist_items:
        row = table.add_row().cells
        if is_header:
            row[0].text = ""
            row[1].text = item
            row[2].text = ""
        else:
            parts = item.split(".", 1)
            row[0].text = parts[0].strip() if len(parts) > 1 else ""
            row[1].text = item
            row[2].text = ""

    doc.add_paragraph()

    # Course Info
    add_heading_para(doc, "COURSE INFORMATION", 1)
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = "Table Grid"
    info_data = [
        ("PROGRAM & SEMESTER", program),
        ("COURSE CODE", course_code),
        ("COURSE NAME", course_name),
        ("COURSE FACULTY", faculty_name),
        ("LTPC", ltpc),
    ]
    for i, (k, v) in enumerate(info_data):
        info_table.rows[i].cells[0].text = k
        info_table.rows[i].cells[1].text = str(v)

    doc.add_paragraph()

    # Standard sections
    add_heading_para(doc, "OUTCOME-BASED CURRICULUM", 1)

    add_heading_para(doc, "1. Vision and Mission of the University", 2)
    add_body(doc, "Vision: To educate students in frontier areas of knowledge enabling them to take up challenges as ethical and responsible global citizens.")
    add_body(doc, "Mission: To impart outcome based holistic education and produce globally competitive, ethical and socially responsible human resources.")
    doc.add_paragraph()

    add_heading_para(doc, "2. Vision and Mission of the School", 2)
    add_body(doc, "Vision: To build a future where education and innovation, empowered by collaboration, create a sustainable global impact.")
    add_body(doc, "Mission: Provide rigorous industry-aligned education, enhance skills, foster research and innovation, build industry partnerships, and nurture ethical engineers.")
    doc.add_paragraph()

    add_heading_para(doc, "3. Vision and Mission of the Department", 2)
    add_body(doc, "Vision: To be a role model computer science department imparting research based multidisciplinary competencies.")
    add_body(doc, "Mission: Build student competency, deploy employability skills, foster research and innovation, build industry partnerships, produce ethical engineers.")
    doc.add_paragraph()

    add_heading_para(doc, "4. GA's, PEOs, POs and PSOs of the Program", 2)
    add_placeholder(doc, "[Copy standard GAs, PEOs, POs and PSOs of the program here]")
    doc.add_paragraph()

    add_heading_para(doc, "5. List of KSA Profiles of the Program", 2)
    add_placeholder(doc, "[Copy standard Knowledge, Skill and Attitude profiles here]")
    doc.add_paragraph()

    add_heading_para(doc, "6. Academic Calendar", 2)
    add_placeholder(doc, "[Attach / paste the Academic Calendar for the current session]")
    doc.add_paragraph()

    # Course Design
    add_heading_para(doc, "OUTCOME BASED COURSE DESIGN", 1)
    add_heading_para(doc, f"Course Code: {course_code}    Course Title: {course_name}", 2)
    add_body(doc, f"Course Type: Program Core Course (PCC)    L-T-P-C: {ltpc}")
    doc.add_paragraph()

    add_heading_para(doc, "7. Course Design of Theory component", 2)
    add_heading_para(doc, "Course Objective", 3)
    add_placeholder(doc, "[Write course objective based on your syllabus]")
    doc.add_paragraph()

    add_heading_para(doc, "Course Outcomes (COs)", 3)
    add_placeholder(doc, "[List 4 Course Outcomes with Bloom's level]")
    doc.add_paragraph()

    add_heading_para(doc, "Syllabus / Course Contents", 3)
    if syllabus_text and syllabus_text.strip():
        for para in syllabus_text.strip().split("\n"):
            para = para.strip()
            if para:
                add_body(doc, para)
    else:
        add_placeholder(doc, "[Paste your detailed syllabus / course contents here]")
    doc.add_paragraph()

    add_heading_para(doc, "Textbooks / Reference Books / Online Resources / MOOCs", 3)
    add_placeholder(doc, "[List textbooks, reference books, online resources and MOOCs]")
    doc.add_paragraph()

    add_heading_para(doc, "Lesson Plan", 3)
    add_placeholder(doc, "[Prepare week-wise / lecture-wise lesson plan with CO mapping, pedagogy and ICT tools]")
    doc.add_paragraph()

    add_heading_para(doc, "CO-PO Mapping and CO-KSA Mapping", 3)
    add_placeholder(doc, "[Add CO-PO and CO-KSA mapping tables]")
    doc.add_paragraph()

    add_heading_para(doc, "Course Assessment Plan", 3)
    add_placeholder(doc, "[Add assessment plan with weightages]")
    doc.add_paragraph()

    add_heading_para(doc, "8. Course Design of Lab component", 2)
    add_placeholder(doc, "[List of experiments, lab COs, rubrics, mini-projects]")
    doc.add_paragraph()

    # Teaching & Learning
    add_heading_para(doc, "OUTCOME BASED TEACHING & LEARNING", 1)
    for title in [
        "9. Setting CO Targets",
        "10. List of faculty members teaching the course",
        "11. Time table of concerned faculty members",
        "12. Lecture Notes and PowerPoint handouts",
        "13. Tutorials/Assignments",
        "14. Solution of Tutorials/Assignments",
        "15. Mid-Sem Question Paper",
        "16. End-Term Question Paper",
        "17. Solution of Mid-Sem and End-Term Question Papers",
        "18. Question Papers / Continuous Assessment",
        "19. Rubrics for Lab assessments",
        "20. Note on ePortfolio of Students",
        "21. Attendance Record for theory and labs",
    ]:
        add_heading_para(doc, title, 2)
        add_placeholder(doc)
        doc.add_paragraph()

    # Assessment
    add_heading_para(doc, "OUTCOME-BASED ASSESSMENT & ATTAINMENT", 1)
    for title in [
        "22. Continuous Evaluation from ERP",
        "23. Assessment-wise Marks Uploaded on ERP",
        "24. Records of Rubrics based Evaluation",
        "25. Records of Assessments based on ePortfolios",
        "26. Minutes of Course Coordinator Level Meetings",
        "27. CO attainment analysis and action taken",
        "28. List of slow Learners and Advanced Learners",
        "29. Action taken to support slow and advanced learners",
        "30. Impact Analysis Report and Action Taken",
        "31. Alternative Plan of Action",
    ]:
        add_heading_para(doc, title, 2)
        add_placeholder(doc)
        doc.add_paragraph()

    # Appendices
    add_heading_para(doc, "APPENDICES", 1)
    add_heading_para(doc, "A. Sample hardcopies of Tutorials and Assignments", 2)
    add_placeholder(doc, "[Attach samples]")
    doc.add_paragraph()
    add_heading_para(doc, "B. Sample hardcopies of answer sheets", 2)
    add_placeholder(doc, "[Attach samples]")

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')} | Developed by aghaZ")
    set_run_font(run, 9)
    run.italic = True

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ====================== UI ======================
st.title("Course File Structure Checker & Generator")

with st.sidebar:
    st.markdown("### Developed by **aghaZ**")
    st.markdown("---")
    st.header("Settings")
    threshold = st.slider("Heading Match Sensitivity", 0.50, 0.90, 0.65, 0.01)
    min_content = st.number_input("Minimum content under heading (chars)", 10, 200, 30, 5)

    st.markdown("---")
    st.header("Dr. Sachin Standard File")
    st.caption("Needed only for Check Structure")

    existing = load_standard_file()
    if existing:
        st.success(f"Loaded: {existing.name}")
        if st.button("Remove Standard"):
            try:
                existing.unlink()
            except Exception:
                pass
            st.rerun()
    else:
        st.warning("No standard file uploaded yet")

    std_up = st.file_uploader("Upload Dr. Sachin Course File", type=["pdf", "docx", "txt"], key="std")
    if std_up:
        save_uploaded_file(std_up, STANDARD_PATH)
        st.success("Standard file saved")
        st.rerun()

tab1, tab2 = st.tabs(["Check Structure", "Generate Course File"])

# TAB 1
with tab1:
    st.subheader("Check Faculty Course File")
    st.markdown("Upload any faculty course file to see which required headings are missing or blank.")

    your_file = st.file_uploader("Upload faculty course file", type=["pdf", "docx", "txt"], key="faculty")

    if st.button("Check Structure", type="primary", use_container_width=True, key="check_btn"):
        std_path = load_standard_file()
        if not std_path:
            st.error("Please upload Dr. Sachin's standard file from the sidebar first.")
        elif not your_file:
            st.error("Please upload a faculty course file.")
        else:
            try:
                with st.spinner("Checking headings..."):
                    your_text = extract_full_text(your_file, is_upload=True)
                    if not your_text.strip():
                        st.error("Could not extract text from the faculty file.")
                    else:
                        results, coverage, missing_count = check_required_headings(
                            your_text, threshold=threshold, min_content=min_content
                        )
                        c1, c2, c3 = st.columns(3)
                        c1.metric("Total Required Headings", len(REQUIRED_HEADINGS))
                        c2.metric("Missing / Blank", missing_count)
                        c3.metric("Coverage", f"{coverage}%")
                        st.progress(min(coverage / 100, 1.0))

                        st.subheader("Checklist Result")
                        problems = [r for r in results if r["status"] != "Present"]
                        present = [r for r in results if r["status"] == "Present"]

                        if problems:
                            st.error(f"{len(problems)} headings are Missing or have blank content")
                            for r in problems:
                                with st.expander(f"{r['required']} — {r['status']}"):
                                    st.write(f"Reason: {r['reason']}")
                                    if r["matched_line"]:
                                        st.write(f"Closest match: {r['matched_line']} (score {r['score']:.2f})")
                        else:
                            st.success("All required structural headings are present and contain content!")

                        with st.expander("Present Headings"):
                            for r in present:
                                st.markdown(f"- **{r['required']}**")

                        report = generate_report(results, coverage, missing_count,
                                                 your_file.name, std_path.name, threshold)
                        st.download_button(
                            "Download Full Report",
                            data=report,
                            file_name=f"structure_check_{datetime.now().strftime('%Y%m%d_%H%M')}.md",
                            mime="text/markdown",
                            use_container_width=True
                        )
            except Exception as e:
                st.error("An error occurred")
                st.code(str(e))
                with st.expander("Details"):
                    st.code(traceback.format_exc())

# TAB 2
with tab2:
    st.subheader("Generate a New Course File")
    st.markdown("Enter details + upload syllabus. The app creates a complete course file with official structure.")

    col1, col2 = st.columns(2)
    with col1:
        course_code = st.text_input("Course Code *", placeholder="e.g. CSH215C")
        course_name = st.text_input("Subject / Course Name *", placeholder="e.g. Computer Networks")
        faculty_name = st.text_input("Faculty Name *", placeholder="e.g. Dr. Manoj Kumar")
    with col2:
        program = st.text_input("Program & Semester", value="B Tech CSE")
        session = st.text_input("Session", value="Jan-Jun, 2026")
        ltpc = st.text_input("L-T-P-C", value="3-0-2-4")

    syllabus_file = st.file_uploader("Upload your Syllabus (PDF / DOCX / TXT)", type=["pdf", "docx", "txt"], key="syllabus")

    if st.button("Generate Course File", type="primary", use_container_width=True, key="gen_btn"):
        if not course_code or not course_name or not faculty_name:
            st.error("Please fill Course Code, Course Name and Faculty Name.")
        else:
            with st.spinner("Generating course file..."):
                try:
                    syllabus_text = ""
                    if syllabus_file:
                        syllabus_text = extract_full_text(syllabus_file, is_upload=True)

                    buffer = create_course_file(
                        course_code=course_code.strip(),
                        course_name=course_name.strip(),
                        faculty_name=faculty_name.strip(),
                        program=program.strip() or "B Tech CSE",
                        session=session.strip() or "Jan-Jun, 2026",
                        syllabus_text=syllabus_text,
                        ltpc=ltpc.strip() or "3-0-2-4"
                    )

                    safe_name = re.sub(r'[^\w\s-]', '', course_name).strip().replace(' ', '_')
                    file_name = f"CourseFile_{course_code}_{safe_name}.docx"

                    st.success("Course file generated successfully!")
                    st.download_button(
                        label="Download Generated Course File (.docx)",
                        data=buffer,
                        file_name=file_name,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
                    st.info("Sections marked [To be filled by faculty] should be completed by the course instructor.")
                except Exception as e:
                    st.error("Failed to generate course file")
                    st.code(str(e))
                    with st.expander("Details"):
                        st.code(traceback.format_exc())

st.markdown("---")
st.caption("Course File Structure Checker & Generator • Developed by aghaZ")
